# -*- coding: utf-8 -*-
"""MD → DOCX 转换（核心文档 + W1/W2 讲稿样板）"""
import re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor

def md_to_docx(md_path, out_path, template_note=''):
    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(2.5); section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5); section.bottom_margin = Cm(2.5)
    lines = open(md_path, encoding='utf-8').read().split('\n')
    i = 0; in_table = False; table_rows = []
    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows: return
        ncols = max(len(r) for r in table_rows)
        t = doc.add_table(rows=len(table_rows), cols=ncols)
        t.style = 'Table Grid'
        for ri, row in enumerate(table_rows):
            for ci in range(ncols):
                t.cell(ri, ci).text = row[ci] if ci < len(row) else ''
        doc.add_paragraph()
        in_table = False; table_rows = []
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith('|'):
            cells = [c.strip() for c in line.strip('|').split('|')]
            if all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                i += 1; continue
            in_table = True; table_rows.append(cells); i += 1; continue
        elif in_table:
            flush_table()
        hm = re.match(r'^(#{1,4})\s+(.*)', line)
        if hm:
            text = re.sub(r'\*\*([^*]+)\*\*', r'\1', hm.group(2).strip())
            text = re.sub(r'`([^`]+)`', r'\1', text)
            doc.add_heading(text, level=min(len(hm.group(1)), 4)); i += 1; continue
        if line.startswith('>'):
            p = doc.add_paragraph(); r = p.add_run(line.lstrip('> ').strip())
            r.italic = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x59,0x59,0x59)
            i += 1; continue
        if re.match(r'^[-*]\s+', line):
            text = re.sub(r'\*\*([^*]+)\*\*', r'\1', re.sub(r'^[-*]\s+', '', line))
            p = doc.add_paragraph(style='List Bullet'); p.add_run(text); i += 1; continue
        if re.match(r'^\d+[.、]\s+', line):
            text = re.sub(r'\*\*([^*]+)\*\*', r'\1', re.sub(r'^\d+[.、]\s+', '', line))
            p = doc.add_paragraph(style='List Number'); p.add_run(text); i += 1; continue
        if line.startswith('```'):
            i += 1; code_lines = []
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i]); i += 1
            p = doc.add_paragraph(); r = p.add_run('\n'.join(code_lines))
            r.font.name = 'Consolas'; r.font.size = Pt(9); i += 1; continue
        if re.match(r'^-{3,}$', line) or re.match(r'^\*{3,}$', line):
            doc.add_paragraph('─' * 40); i += 1; continue
        if line.strip():
            text = re.sub(r'\*\*([^*]+)\*\*', r'\1', line)
            text = re.sub(r'\*([^*]+)\*', r'\1', text)
            text = re.sub(r'`([^`]+)`', r'\1', text)
            doc.add_paragraph(text)
        i += 1
    if in_table: flush_table()
    if template_note:
        p = doc.add_paragraph(); r = p.add_run(f'【模板说明】{template_note}')
        r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x99,0x99,0x99)
    doc.save(out_path)
    print('OK ' + out_path)

JOBS = [
    ("01_工作调研与课程标准/S04_课程标准/课程标准.md",
     "01_工作调研与课程标准/S04_课程标准/word版/课程标准（默认模板）.docx", "L2-01 课程标准模板"),
    ("03_学情与整体设计/S06_课程整体设计/微电影编导-课程整体设计（三阶梯项目制·情人事版）.md",
     "03_学情与整体设计/S06_课程整体设计/word版/微电影编导-课程整体设计（三阶梯项目制·情人事版）（默认模板）.docx", "L2-02 课程整体设计模板"),
    ("08_授课计划/授课计划-微电影编导.md",
     "08_授课计划/word版/授课计划-微电影编导（默认模板）.docx", "L2-04 授课计划模板"),
    ("04_单元教学设计/单元设计总览_微电影编导_13周.md",
     "04_单元教学设计/word版/单元设计总览_微电影编导_13周（默认模板）.docx", "L2-03 单元设计模板"),
    ("05_课堂备课包/备课包_第1周_作品P1_情境认知与画面思维/05_讲稿/讲稿_W1_情境认知.md",
     "05_课堂备课包/word版/讲稿_W1_情境认知（默认模板）.docx", "L3-05 讲稿模板"),
    ("05_课堂备课包/备课包_第2周_作品P1_无对白创作与分镜/05_讲稿/讲稿_W2_无对白创作.md",
     "05_课堂备课包/word版/讲稿_W2_无对白创作（默认模板）.docx", "L3-05 讲稿模板"),
]

for md, out, note in JOBS:
    if not os.path.exists(md):
        print('MISS ' + md); continue
    try:
        md_to_docx(md, out, note)
    except Exception as e:
        print('ERR ' + md + ': ' + str(e))
print('ALL DONE')
