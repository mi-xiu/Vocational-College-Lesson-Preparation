# -*- coding: utf-8 -*-
"""单元教案通道 · 通用段落引擎（模板导出专家 · Ti08）

可复用原语（课程无关，纯模板操作）：
- 抓格式：get_rpr / get_ppr
- 改格式：force_rpr（强制某段为节标题粗体 rPr）
- 改文本：set_para_text（保留首 run 原 rPr）
- 建段落：make_para / make_para_mixed —— 新插入段落**必须同时继承 pPr + rPr**
- 删块：delete_block_between（删除过程设计块，保留前后章节）
- 清扫：sweep_replace（课程换皮安全网）
- 定位：locate_by_keyword / locate_after_heading（过程块删除后索引漂移，用文本定位）

关键纪律（已固化）：
新插入段落只复制 run 的 rPr（字体）会丢失段落 pPr（首行缩进 ind firstLineChars、
行距 spacing line），导致新段回退为无缩进/默认行距。**必须同时复制 pPr**。
"""
import sys, os
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 复用 school-template 渲染器的保真底层函数（replace_in_para / replace_all）
HERE = os.path.dirname(os.path.abspath(__file__))
ST = os.path.join(os.path.dirname(HERE), "school-template")
if ST not in sys.path:
    sys.path.insert(0, ST)
import render_template as RT


def get_rpr(paragraph):
    for r in paragraph.runs:
        rpr = r._element.find(qn('w:rPr'))
        if rpr is not None:
            return deepcopy(rpr)
    return None


def get_ppr(paragraph):
    pPr = paragraph._element.find(qn('w:pPr'))
    return deepcopy(pPr) if pPr is not None else None


def force_rpr(paragraph, rpr):
    """强制把某段所有 run 的 rPr 替换为模板节标题的粗体 rPr。"""
    for r in paragraph.runs:
        old = r._element.find(qn('w:rPr'))
        if old is not None:
            r._element.remove(old)
        if rpr is not None:
            r._element.insert(0, deepcopy(rpr))


def set_para_text(paragraph, new_text):
    """整段替换文本，保留首 run 的 rPr（与 RT.replace_in_para 行为一致）。"""
    full = "".join(r.text for r in paragraph.runs) if paragraph.runs else paragraph.text
    RT.replace_in_para(paragraph, full, new_text)


def make_para(text, rpr, ppr=None):
    """新建单行段落：可带节标题粗体 rPr + 正文 pPr。"""
    p = OxmlElement('w:p')
    if ppr is not None:
        p.append(deepcopy(ppr))
    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(deepcopy(rpr))
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    p.append(r)
    return p


def make_para_mixed(prefix, content, rpr_bold, rpr_norm, ppr=None):
    """新建混合段落：前缀粗体（环节名）+ 正文常规，整段继承 pPr。"""
    p = OxmlElement('w:p')
    if ppr is not None:
        p.append(deepcopy(ppr))
    r1 = OxmlElement('w:r')
    if rpr_bold is not None:
        r1.append(deepcopy(rpr_bold))
    t1 = OxmlElement('w:t')
    t1.text = prefix
    t1.set(qn('xml:space'), 'preserve')
    r1.append(t1)
    p.append(r1)
    r2 = OxmlElement('w:r')
    if rpr_norm is not None:
        r2.append(deepcopy(rpr_norm))
    t2 = OxmlElement('w:t')
    t2.text = content
    t2.set(qn('xml:space'), 'preserve')
    r2.append(t2)
    p.append(r2)
    return p


def delete_block_between(doc, anchor_idx, end_marker):
    """删除 anchor_idx 之后、直到文本以 end_marker 开头的段落之前的所有段落。"""
    p_anchor = doc.paragraphs[anchor_idx]._element
    nxt = p_anchor.getnext()
    while nxt is not None:
        txt = "".join(t.text or "" for t in nxt.iter(qn('w:t')))
        if txt.strip().startswith(end_marker):
            break
        n = nxt.getnext()
        nxt.getparent().remove(nxt)
        nxt = n


def sweep_replace(doc, pairs):
    """课程换皮安全网：模板残留旧文本统一替换。"""
    for old, new in pairs:
        RT.replace_all(doc, old, new)


def locate_by_keyword(doc, keyword, exclude_keyword=None):
    """按关键字定位首个匹配段落（过程块删除后索引漂移，用文本定位最稳）。"""
    for p in doc.paragraphs:
        txt = "".join(r.text for r in p.runs) if p.runs else p.text
        if keyword in txt and (exclude_keyword is None or exclude_keyword not in txt):
            return p
    return None


def locate_after_heading(doc, heading_text):
    """返回文本精确等于 heading_text 的段落的下一段落（用于"作业"等小标题后的正文）。"""
    for i, p in enumerate(doc.paragraphs):
        txt = "".join(r.text for r in p.runs) if p.runs else p.text
        if txt.strip() == heading_text:
            return doc.paragraphs[i + 1]
    return None
