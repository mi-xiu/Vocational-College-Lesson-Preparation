# -*- coding: utf-8 -*-
"""模板优化器：给定学校模板 docx，抽取"可执行渲染规格"写入 spec.json 的 render 段，
并产出数据骨架 <stem>-data.example.json。

这是"先做优化"的一步——把原来要手写脚本探查的合并结构、列宽、合计行一次性固化进规格，
后续同类课程只需填数据 json + 跑 render_template.py，无需再写 Python。

用法：
  python optimize_template.py <模板.docx> [<现有spec.json可选>]
"""
import sys, os, re, json
from docx import Document
from docx.oxml.ns import qn

def cell_info(cell):
    tcPr = cell._tc.find(qn('w:tcPr'))
    gs, vm, w = "", "", ""
    if tcPr is not None:
        g = tcPr.find(qn('w:gridSpan'))
        if g is not None:
            gs = g.get(qn('w:val'))
        v = tcPr.find(qn('w:vMerge'))
        if v is not None:
            vm = v.get(qn('w:val')) or 'restart'
        we = tcPr.find(qn('w:tcW'))
        if we is not None:
            w = we.get(qn('w:w'))
    return gs, vm, w

def detect_tables(doc):
    tables = []
    for ti, tbl in enumerate(doc.tables):
        nrows = len(tbl.rows)
        ncols = len(tbl.columns)
        header = [c.text for c in tbl.rows[0].cells]
        # 合并列 & 跨行 span 检测
        merge_cols = set()
        has_span = False  # 存在跨多行的合并（restart 后接 continue）
        for ci in range(ncols):
            for ri in range(1, nrows):
                gs, vm, _ = cell_info(tbl.rows[ri].cells[ci])
                if vm:
                    merge_cols.add(ci)
                    if vm == 'restart' and ri + 1 < nrows:
                        g2, vm2, _ = cell_info(tbl.rows[ri + 1].cells[ci])
                        if vm2 == 'continue':
                            has_span = True
        # 列宽（取 row1）
        widths = []
        ref = tbl.rows[1] if nrows > 1 else tbl.rows[0]
        for c in ref.cells:
            _, _, w = cell_info(c)
            widths.append(w)
        # 合计行检测
        total = False
        total_label = {}
        last = tbl.rows[-1]
        for ci, c in enumerate(last.cells):
            gs, vm, _ = cell_info(c)
            if c.text.strip() == "合计" or (gs and int(gs) > 1) or vm == 'restart':
                total = True
            if c.text.strip():
                total_label[str(ci)] = c.text.strip()
        # 表头行数：连续前导行且每行所有单元格都带 vm restart
        header_rows = 0
        for ri in range(nrows):
            if all(cell_info(c)[1] == 'restart' for c in tbl.rows[ri].cells):
                header_rows += 1
            else:
                break
        if header_rows == 0:
            header_rows = 1
        # 模式决策
        if has_span:
            mode = "cells"
        elif merge_cols:
            mode = "rebuild"
        else:
            mode = "fill"
        entry = {"idx": ti, "mode": mode, "header_rows": header_rows}
        if mode == "rebuild":
            entry["merge_cols"] = sorted(merge_cols)
            entry["header"] = header
            entry["col_widths"] = widths
        elif mode == "fill":
            if total:
                entry["total"] = True
                entry["total_label"] = total_label
        tables.append(entry)
    return tables

def detect_identity(doc):
    """探测可清扫的模板专属标记（课程名/专业/代码/学时/日期）。"""
    cover_title = ""
    major = ""
    code = ""
    hours = ""
    date = ""
    # 封面大标题：前 12 个非空段落里字号最大且加粗者
    cands = []
    for p in doc.paragraphs[:40]:
        if not p.text.strip():
            continue
        sz = 0
        b = False
        for r in p.runs:
            rpr = r._element.find(qn('w:rPr'))
            if rpr is not None:
                sz_e = rpr.find(qn('w:sz'))
                if sz_e is not None:
                    sz = max(sz, int(sz_e.get(qn('w:val'))))
                if rpr.find(qn('w:b')) is not None:
                    b = True
        cands.append((sz, b, p.text.strip()))
        if len(cands) >= 12:
            break
    if cands:
        cover_title = max(cands, key=lambda x: (x[0], x[1]))[2]
    full = "\n".join(p.text for p in doc.paragraphs)
    m = re.search(r"适用专业[:：]\s*([^\n]+?)(?:\s|$)", full)
    if not m:
        m = re.search(r"([^\n]+?)(?:专业教研室|教研室\s*制)", full)
    if m:
        major = m.group(1).strip().split()[-1] if m.group(1) else ""
        # 取包含"专业"的整词
        mm = re.search(r"([\u4e00-\u9fa5]{2,}专业)", m.group(1))
        if mm:
            major = mm.group(1)
    m = re.search(r"课程代码[:：]\s*([0-9]+)", full)
    if m:
        code = m.group(1)
    m = re.search(r"(\d+\s*学时)", full)
    if m:
        hours = m.group(1)
    m = re.search(r"(\d{4}\.\d{1,2}\.\d{1,2})", full)
    if m:
        date = m.group(1)
    sweep = []
    for tok in [cover_title, major, code, hours, date]:
        if tok and tok not in [s[0] for s in sweep]:
            sweep.append(tok)
    return sweep

def build_skeleton(doc, tables, sweep):
    paragraphs = {}
    for pi, p in enumerate(doc.paragraphs):
        if p.text.strip():
            paragraphs[str(pi)] = p.text
    tabs = {}
    totals = {}
    cells_values = {}
    for te in tables:
        ti = te["idx"]
        tbl = doc.tables[ti]
        if te["mode"] == "cells":
            cv = {}
            for ri, row in enumerate(tbl.rows):
                for ci, c in enumerate(row.cells):
                    if c.text.strip():
                        cv[f"c{ri}_{ci}"] = c.text
            cells_values[str(ti)] = cv
        else:
            rows = []
            for ri in range(te.get("header_rows", 1), len(tbl.rows)):
                # 合计行跳过（由 total_label 处理）
                if te.get("mode") == "fill" and te.get("total") and ri == len(tbl.rows) - 1:
                    continue
                rows.append([c.text for c in tbl.rows[ri].cells])
            tabs[str(ti)] = rows
            if te.get("mode") == "fill" and te.get("total"):
                totals[str(ti)] = {}
    sweep_pairs = [[s, ""] for s in sweep if s]
    return {
        "paragraphs": paragraphs,
        "tables": tabs,
        "totals": totals,
        "cells_values": cells_values,
        "sweep": sweep_pairs,
    }

def main():
    tpl_path = sys.argv[1]
    spec_path = sys.argv[2] if len(sys.argv) > 2 else None
    doc = Document(tpl_path)
    tables = detect_tables(doc)
    sweep = detect_identity(doc)
    render_block = {"paragraphs": [], "tables": tables}
    # 段落：所有非空段落登记为可填充字段
    for pi, p in enumerate(doc.paragraphs):
        if p.text.strip():
            render_block["paragraphs"].append({
                "idx": pi, "match": p.text.strip(),
                "hint": p.text.strip()[:12],
            })
    # 写回 spec
    spec = {}
    if spec_path and os.path.exists(spec_path):
        with open(spec_path, encoding="utf-8") as f:
            spec = json.load(f)
    spec["source_file"] = os.path.basename(tpl_path)
    spec["render"] = render_block
    out_spec = spec_path or (os.path.splitext(tpl_path)[0] + "-spec.json")
    with open(out_spec, "w", encoding="utf-8") as f:
        json.dump(spec, f, ensure_ascii=False, indent=1)
    # 数据骨架
    stem = os.path.splitext(tpl_path)[0]
    skel = build_skeleton(doc, tables, sweep)
    skel_path = stem + "-data.example.json"
    with open(skel_path, "w", encoding="utf-8") as f:
        json.dump(skel, f, ensure_ascii=False, indent=1)
    print("优化完成：")
    print("  规格 ->", out_spec)
    print("  数据骨架 ->", skel_path)
    print("  表格模式：", [(t['idx'], t['mode']) for t in tables])
    print("  清扫标记：", sweep)

if __name__ == "__main__":
    main()
