# -*- coding: utf-8 -*-
"""《微电影编导》课程 PPT 生成器（第 1 周样板 · 20 页）
风格：现代清爽·高端教育科技（蓝白主色 + 青蓝/暖橙/珊瑚点缀 + 渐变光晕）
方法：吸收 ppt-master 方法论——原生形状优先、大师级信息层级、16:9 画布
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---- 主色板 ----
BLUE = RGBColor(0x25, 0x63, 0xEB)       # 主蓝
DARKBLUE = RGBColor(0x1E, 0x3A, 0x8A)   # 深蓝
CYAN = RGBColor(0x06, 0xB6, 0xD4)       # 亮青蓝
ORANGE = RGBColor(0xF5, 0x9E, 0x0B)     # 暖橙
CORAL = RGBColor(0xFB, 0x71, 0x85)      # 珊瑚
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHTBLUE = RGBColor(0xEF, 0xF6, 0xFF)  # 浅蓝底
GREY = RGBColor(0x47, 0x55, 0x69)       # 灰文字
LIGHTGREY = RGBColor(0x94, 0xA3, 0xB8)  # 浅灰

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)  # 16:9

def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

def add_rect(slide, x, y, w, h, color, shape_type=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape_type, x, y, w, h)
    set_fill(sh, color)
    return sh

def add_rounded(slide, x, y, w, h, color, radius=0.08):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    set_fill(sh, color)
    return sh

def add_glow(slide, x, y, w, h, color, alpha=0.15):
    """半透明光晕椭圆（模拟柔和光晕）"""
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    # 设置透明度
    solidFill = sh.fill._xPr.find(qn('a:solidFill'))
    if solidFill is None:
        solidFill = sh.fill._xPr.makeelement(qn('a:solidFill'), {})
        sh.fill._xPr.append(solidFill)
    srgb = solidFill.find(qn('a:srgbClr'))
    if srgb is not None:
        alpha_elem = srgb.makeelement(qn('a:alpha'), {'val': str(int(alpha * 100000))})
        srgb.append(alpha_elem)
    sh.line.fill.background()
    return sh

def add_text(slide, x, y, w, h, text, size=18, color=GREY, bold=False,
             align=PP_ALIGN.LEFT, font='微软雅黑', anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(4)
    tf.margin_top = tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    # 设置中文字体
    rPr = run._r.get_or_add_rPr()
    ea = rPr.makeelement(qn('a:ea'), {'typeface': font})
    rPr.append(ea)
    return tb

def add_card(slide, x, y, w, h, title, body, title_color=BLUE, bg=LIGHTBLUE, tsize=16, bsize=12):
    """要点卡片：圆角矩形 + 标题 + 正文"""
    card = add_rounded(slide, x, y, w, h, bg)
    card.line.color.rgb = RGBColor(0xDB, 0xE6, 0xFF)
    card.line.width = Pt(1)
    # 标题
    add_text(slide, x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), Inches(0.32),
             title, size=tsize, color=title_color, bold=True)
    # 正文
    add_text(slide, x + Inches(0.15), y + Inches(0.42), w - Inches(0.3), h - Inches(0.52),
             body, size=bsize, color=GREY)

def add_gradient_bg(slide, c1=RGBColor(0xEF, 0xF6, 0xFF), c2=RGBColor(0xF0, 0xFD, 0xFF)):
    """渐变背景（模拟光滑柔和渐变）"""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.line.fill.background()
    fill = bg.fill
    fill.gradient()
    stops = fill.gradient_stops
    if len(stops) >= 2:
        stops[0].color.rgb = c1
        stops[1].color.rgb = c2
    bg.shadow.inherit = False
    return bg

def add_footer(slide, page_no, text='微电影编导 · 第 1 周 · 情境认知与画面思维'):
    """页脚：页码 + 课程名"""
    add_text(slide, Inches(0.5), Inches(7.05), Inches(8), Inches(0.3),
             text, size=9, color=LIGHTGREY)
    add_text(slide, Inches(12.5), Inches(7.05), Inches(0.6), Inches(0.3),
             str(page_no), size=10, color=LIGHTGREY, align=PP_ALIGN.RIGHT)

def add_title_bar(slide, title, subtitle='', accent=CYAN):
    """统一标题栏：色条 + 大标题 + 副标题"""
    add_rect(slide, Inches(0.6), Inches(0.55), Inches(0.08), Inches(0.6), accent)
    add_text(slide, Inches(0.85), Inches(0.5), Inches(11), Inches(0.6),
             title, size=28, color=DARKBLUE, bold=True)
    if subtitle:
        add_text(slide, Inches(0.85), Inches(1.05), Inches(11), Inches(0.35),
                 subtitle, size=13, color=LIGHTGREY)

def add_building_skyline(slide, y=Inches(6.3), color=RGBColor(0xDB, 0xE6, 0xFF), alpha=0.4):
    """远处校园建筑轮廓（简洁层次，边缘柔和）"""
    buildings = [(0.8, 0.35), (1.4, 0.55), (2.0, 0.4), (2.7, 0.6), (3.5, 0.45),
                 (4.2, 0.58), (5.0, 0.38), (5.8, 0.52), (6.5, 0.42), (7.3, 0.55),
                 (8.1, 0.4), (8.9, 0.5), (9.7, 0.36), (10.5, 0.53), (11.3, 0.45), (12.0, 0.5)]
    for bx, bw in buildings:
        bh = Inches(0.5 + bw * 0.3)
        sh = add_rect(slide, Inches(bx), y - bh, Inches(bw * 1.2), bh, color)
        # 透明度
        solidFill = sh.fill._xPr.find(qn('a:solidFill'))
        if solidFill is not None:
            srgb = solidFill.find(qn('a:srgbClr'))
            if srgb is not None:
                a = srgb.makeelement(qn('a:alpha'), {'val': str(int(alpha * 100000))})
                srgb.append(a)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank = prs.slide_layouts[6]

# ============ P1 封面 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s, RGBColor(0xDB, 0xEA, 0xFF), RGBColor(0xE0, 0xF7, 0xFF))
# 光晕
add_glow(s, Inches(9.5), Inches(-1), Inches(5), Inches(5), CYAN, alpha=0.12)
add_glow(s, Inches(-1.5), Inches(4.5), Inches(4), Inches(4), ORANGE, alpha=0.08)
# 顶部标签
add_rounded(s, Inches(0.9), Inches(0.8), Inches(2.2), Inches(0.42), BLUE, radius=0.5)
add_text(s, Inches(1.1), Inches(0.86), Inches(1.9), Inches(0.3),
         '微电影编导 · 第 1 周', size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
# 主标题
add_text(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.2),
         '情境认知与画面思维', size=54, color=DARKBLUE, bold=True)
add_text(s, Inches(0.9), Inches(3.2), Inches(11), Inches(0.6),
         '作品 P1 · 诗意小品 · 以人物动作讲情绪，不给剧本，你来创作', size=20, color=BLUE)
# 副标签
add_text(s, Inches(0.9), Inches(4.1), Inches(8), Inches(0.4),
         '教学情境命题：等待 / 归途 / 邂逅 —— 西溪湿地意象', size=14, color=GREY)
# 底部信息条
add_rounded(s, Inches(0.9), Inches(5.3), Inches(5.5), Inches(0.55), LIGHTBLUE, radius=0.3)
add_text(s, Inches(1.1), Inches(5.38), Inches(5.1), Inches(0.4),
         'BOtPPPS · 180 分钟 · 量规前置 ★闸门', size=12, color=DARKBLUE, bold=True)
# 建筑轮廓
add_building_skyline(s, Inches(6.4), RGBColor(0xC7, 0xDB, 0xFF), alpha=0.35)
add_footer(s, 1)

# ============ P2 今日任务 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '今日任务', '交付物明确 · 量规前置', CYAN)
add_card(s, Inches(0.7), Inches(1.7), Inches(5.8), Inches(2.0),
         '① 情境理解笔记', '选情境（等待/归途/邂逅）→ 情绪方向 → 一句话故事\n→ 无对白题材判断（删掉台词故事还成立吗）', BLUE)
add_card(s, Inches(6.9), Inches(1.7), Inches(5.8), Inches(2.0),
         '② 拉片入门笔记', '《一无所有》《两个人的沉默》→ 动作/状态/关系\n在讲什么情绪（教材①第三章范例）', CYAN)
add_card(s, Inches(0.7), Inches(4.0), Inches(5.8), Inches(1.6),
         '预告 W2', '自创无对白剧本（不给剧本！）——画面点子 = 剧本骨架', ORANGE)
add_card(s, Inches(6.9), Inches(4.0), Inches(5.8), Inches(1.6),
         '★ 量规前置', '达成度 ≥80% 且 ★ 全过 → 通过；未过返工复评', CORAL)
add_footer(s, 2)

# ============ P3 情境命题卡 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, 'P1 情境命题卡', '三选一 · 不给剧本 · 你来创作', ORANGE)
# 三个选项卡片
opts = [
    ('A. 等待', '芦苇荡码头 · 一个人望着水面\n握着信 · 守望与期盼', BLUE),
    ('B. 归途', '河渚街黄昏 · 一个人走在老街\n停下回头 · 乡愁与归家', CYAN),
    ('C. 邂逅', '摇橹船码头 · 两个人擦肩对视\n错身 · 缘分与错过', CORAL),
]
for i, (t, d, c) in enumerate(opts):
    x = Inches(0.7 + i * 4.2)
    add_card(s, x, Inches(1.8), Inches(3.9), Inches(3.2), t, d, c, tsize=20, bsize=14)
add_text(s, Inches(0.7), Inches(5.5), Inches(12), Inches(0.5),
         '★ 创作要求：人物动作/状态/关系叙事为主，物为氛围非主角；30-60 秒无对白', size=14, color=DARKBLUE, bold=True)
add_footer(s, 3)

# ============ P4 西溪湿地意象 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s, RGBColor(0xE0, 0xF7, 0xFF), RGBColor(0xF0, 0xFD, 0xFF))
add_title_bar(s, '西溪湿地 · 意象池', '芦苇荡 / 河渚街 / 摇橹船码头', CYAN)
# 意象卡片（模拟空镜位）
imgs = [
    ('芦苇荡码头', '芦花摇曳 · 水面波光', '等待 · 守望'),
    ('河渚街黄昏', '暖光斜照 · 老街巷尾', '归途 · 乡愁'),
    ('摇橹船码头', '船影轻晃 · 人来人往', '邂逅 · 缘分'),
]
for i, (t, d, e) in enumerate(imgs):
    x = Inches(0.7 + i * 4.2)
    add_rounded(s, x, Inches(1.8), Inches(3.9), Inches(2.6), RGBColor(0xD5, 0xE8, 0xFF), radius=0.06)
    add_text(s, x + Inches(0.2), Inches(2.0), Inches(3.5), Inches(0.5),
             t, size=18, color=DARKBLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.7), Inches(3.5), Inches(0.4),
             d, size=13, color=GREY, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(3.6), Inches(3.5), Inches(0.4),
             '→ ' + e, size=14, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0.7), Inches(4.9), Inches(12), Inches(0.5),
         '媒体：西溪意象素材（30 秒）｜缺口未补则用 07_动画 情境卡', size=12, color=LIGHTGREY)
add_footer(s, 4)

# ============ P5 导入提问 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '导入提问', '先说出你的直觉', CORAL)
add_glow(s, Inches(9), Inches(2.5), Inches(3.5), Inches(3.5), CORAL, alpha=0.08)
add_text(s, Inches(0.7), Inches(2.2), Inches(11.5), Inches(1.2),
         '"这个画面在讲什么情绪？"', size=36, color=DARKBLUE, bold=True)
add_text(s, Inches(0.7), Inches(3.5), Inches(11.5), Inches(0.8),
         '你是看什么判断的——人，还是景？', size=24, color=BLUE)
add_card(s, Inches(0.7), Inches(4.7), Inches(11.9), Inches(1.4),
         '思政 · 情境美育', '生态之美 / 乡愁 / 缘分——画面先于语言抵达情感', ORANGE, tsize=15)
add_footer(s, 5)

# ============ P6 学习目标 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '学习目标（PGSD）', '本单元能力点 · 原样引用 S03', BLUE)
goals = [
    ('复述', '剧本 = 用画面讲故事（S-05）', BLUE),
    ('识别', '人物动作叙事（G-01）', CYAN),
    ('判断', '情境情绪方向（G-02）', ORANGE),
    ('原创', '自创第一课 · AI 出量人出质（D-02）', CORAL),
]
for i, (v, d, c) in enumerate(goals):
    x = Inches(0.7 + i * 3.1)
    add_rounded(s, x, Inches(1.8), Inches(2.8), Inches(2.8), LIGHTBLUE, radius=0.1)
    add_rect(s, x, Inches(1.8), Inches(2.8), Inches(0.1), c)
    add_text(s, x + Inches(0.2), Inches(2.1), Inches(2.4), Inches(0.5),
             v, size=18, color=c, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.8), Inches(2.4), Inches(1.5),
             d, size=13, color=DARKBLUE, align=PP_ALIGN.CENTER)
add_text(s, Inches(0.7), Inches(5.1), Inches(12), Inches(0.5),
         '对齐教案板块 3 目标｜达成判定：见板块 7 量规（✔/✘ + ★）', size=12, color=LIGHTGREY)
add_footer(s, 6)

# ============ P7 前测 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '前测 · 激活旧知', 'KWL 表 + 情境情绪填空', CYAN)
add_card(s, Inches(0.7), Inches(1.7), Inches(5.8), Inches(3.6),
         '情境情绪填空', '等待 → 情绪是 ____\n画面会有 ____\n（凭直觉填，不评判对错）', BLUE, tsize=16, bsize=15)
add_card(s, Inches(6.9), Inches(1.7), Inches(5.8), Inches(3.6),
         'KWL 表', 'K 已知：我懂哪些画面语言\nW 想知：我想学什么\nL 已学：课后回填', CYAN, tsize=16, bsize=15)
add_text(s, Inches(0.7), Inches(5.6), Inches(12), Inches(0.5),
         '暴露"会做不会说"的起点，为归纳式教学提供对比基线', size=13, color=GREY)
add_footer(s, 7)

# ============ P8 剧本=用画面讲故事 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '剧本 = 用画面讲故事', '电影思维 vs 小说思维', BLUE)
add_card(s, Inches(0.7), Inches(1.7), Inches(5.8), Inches(3.4),
         '小说用句子', '心理描写：\n"他内心充满了等待的焦虑…"\n（文字直达内心）', GREY, tsize=15, bsize=13)
add_card(s, Inches(6.9), Inches(1.7), Inches(5.8), Inches(3.4),
         '电影用画面', '动作描写：\n码头 + 折信纸 → 望水面 → 低头看表\n（行为即心理，观众自己感受）', BLUE, tsize=15, bsize=13)
add_text(s, Inches(0.7), Inches(5.4), Inches(12), Inches(0.6),
         '教材①第一章：同一段"等待"，小说写心理 vs 电影拍动作', size=13, color=ORANGE, bold=True)
add_footer(s, 8)

# ============ P9 一场戏/动作单元 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '一场戏 · 动作单元', '场景 + 人物动作 + 情绪', CYAN)
add_card(s, Inches(0.7), Inches(1.7), Inches(5.8), Inches(2.6),
         '一场戏 = ', '场景（码头）+ 人物动作（折信/望水/看表）\n+ 情绪（守望）——画面点子 ≈ 动作单元', BLUE, tsize=15, bsize=13)
add_card(s, Inches(6.9), Inches(1.7), Inches(5.8), Inches(2.6),
         '板书拆解', '码头 + 折信纸 → 望水面 → 低头看表 = 守望\n（一个动作单元 = 一个画面点子）', CYAN, tsize=15, bsize=13)
add_text(s, Inches(0.7), Inches(4.7), Inches(12), Inches(0.5),
         '教材①第二章：认识"一场戏"——麻雀虽小五脏俱全', size=13, color=GREY)
add_footer(s, 9)

# ============ P10 拉片示范1 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s, RGBColor(0xF0, 0xFD, 0xFF), RGBColor(0xEF, 0xF6, 0xFF))
add_title_bar(s, '拉片示范 1：《一无所有》', '动作 + 关系在讲"孤独"', ORANGE)
add_rounded(s, Inches(0.7), Inches(1.7), Inches(11.9), Inches(3.4), RGBColor(0xE2, 0xF0, 0xFF), radius=0.05)
add_text(s, Inches(1.0), Inches(2.0), Inches(11), Inches(0.5),
         '街头歌手反复拨弦 · 人流走过不停 · 歌声在嘈杂都市里如"溺水呼救"', size=17, color=DARKBLUE, bold=True)
add_text(s, Inches(1.0), Inches(3.0), Inches(11), Inches(0.8),
         '画面：长焦俯拍人流 → 低角度仰拍歌手 → 红男绿女面无表情地走过', size=14, color=GREY)
add_text(s, Inches(1.0), Inches(4.2), Inches(11), Inches(0.6),
         '→ 动作+关系讲"孤独"：无人驻足 = 情感缺席（媒体：教材①第三章范例）', size=14, color=ORANGE, bold=True)
add_footer(s, 10)

# ============ P11 拉片示范2 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s, RGBColor(0xEF, 0xF6, 0xFF), RGBColor(0xF0, 0xFD, 0xFF))
add_title_bar(s, '拉片示范 2：《两个人的沉默》', '状态 + 关系在讲"张力"', CYAN)
add_rounded(s, Inches(0.7), Inches(1.7), Inches(11.9), Inches(3.4), RGBColor(0xE0, 0xF7, 0xFF), radius=0.05)
add_text(s, Inches(1.0), Inches(2.0), Inches(11), Inches(0.5),
         '对坐 · 眼神避开 · 谁都不先开口', size=17, color=DARKBLUE, bold=True)
add_text(s, Inches(1.0), Inches(3.0), Inches(11), Inches(0.8),
         '沉默本身在说话：身体语言 + 空间距离 + 时间流逝 = 关系张力', size=14, color=GREY)
add_text(s, Inches(1.0), Inches(4.2), Inches(11), Inches(0.6),
         '→ 无对白也能"讲"出人与人之间的暗流（媒体：教材①第三章范例）', size=14, color=CYAN, bold=True)
add_footer(s, 11)

# ============ P12 人物动作叙事 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '人物动作叙事', '动作 / 状态 / 关系 讲情绪', BLUE)
acts = [
    ('动作', '折信纸/看表/回头\n行为即心理', BLUE),
    ('状态', '站着/坐着/发呆\n静止也叙事', CYAN),
    ('关系', '擦肩/对视/错身\n人物间有戏', ORANGE),
]
for i, (t, d, c) in enumerate(acts):
    x = Inches(0.7 + i * 4.2)
    add_card(s, x, Inches(1.8), Inches(3.9), Inches(3.0), t, d, c, tsize=20, bsize=14)
add_text(s, Inches(0.7), Inches(5.2), Inches(12), Inches(0.6),
         '物为氛围非主角：芦苇/船/信是"情绪的容器"，人是情感主体', size=14, color=ORANGE, bold=True)
add_footer(s, 12)

# ============ P13 无对白题材边界 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '无对白题材边界', '不是任何题材都适合', CORAL)
add_card(s, Inches(0.7), Inches(1.7), Inches(5.8), Inches(3.2),
         '✅ 适合', '等待 / 归途 / 邂逅 / 离别重逢\n动作与状态足以承载情绪', RGBColor(0x10, 0xB9, 0x81), tsize=16, bsize=14)
add_card(s, Inches(6.9), Inches(1.7), Inches(5.8), Inches(3.2),
         '✘ 不适合', '法庭审理 / 恋爱长对话\n必须靠语言交流的题材', CORAL, tsize=16, bsize=14)
add_text(s, Inches(0.7), Inches(5.2), Inches(12), Inches(0.6),
         '口诀：删掉台词，故事还成立吗？（教材①第三章）', size=15, color=DARKBLUE, bold=True)
add_footer(s, 13)

# ============ P14 情境理解笔记模板 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '情境理解笔记 · 模板', '四步完成（工作页板块 B）', CYAN)
steps = [
    ('1', '选情境', '等待/归途/邂逅 三选一', BLUE),
    ('2', '情绪方向', '这个情境想表达什么情绪', CYAN),
    ('3', '画面点子', '动作叙事（人物动作+状态+关系）', ORANGE),
    ('4', '结尾落点', '情绪的落点（守候/释然/错过）', CORAL),
]
for i, (n, t, d, c) in enumerate(steps):
    x = Inches(0.7 + i * 3.1)
    add_rounded(s, x, Inches(1.8), Inches(2.8), Inches(3.4), LIGHTBLUE, radius=0.1)
    add_glow(s, x + Inches(0.1), Inches(1.9), Inches(0.5), Inches(0.5), c, alpha=0.2)
    add_text(s, x + Inches(0.2), Inches(2.0), Inches(0.6), Inches(0.5),
             n, size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.7), Inches(2.4), Inches(0.5),
             t, size=16, color=c, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(3.4), Inches(2.4), Inches(1.5),
             d, size=12, color=GREY, align=PP_ALIGN.CENTER)
add_footer(s, 14)

# ============ P15 AI 联想 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, 'AI 联想：AI 出量、人出质', '原创自创边界', ORANGE)
add_card(s, Inches(0.7), Inches(1.7), Inches(5.8), Inches(2.8),
         'AI 能做什么', '联想画面点子 / 生成情境参考图\n（文生图：码头+芦苇+等待）', CYAN, tsize=15, bsize=13)
add_card(s, Inches(6.9), Inches(1.7), Inches(5.8), Inches(2.8),
         '人必须做什么', '判断 / 选择 / 留痕\n标注"改了什么、为何改"——判断权在人', ORANGE, tsize=15, bsize=13)
add_text(s, Inches(0.7), Inches(4.9), Inches(12), Inches(0.6),
         '★ 原创边界：AI 联想可加分，但情绪判断与最终取舍必须是你自己的', size=14, color=CORAL, bold=True)
add_footer(s, 15)

# ============ P16 互述提纲 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '后测 · 互述提纲', '每人 1 分钟', BLUE)
add_rounded(s, Inches(0.7), Inches(1.8), Inches(11.9), Inches(3.0), LIGHTBLUE, radius=0.06)
add_text(s, Inches(1.0), Inches(2.2), Inches(11), Inches(0.6),
         '说清三件事：', size=18, color=DARKBLUE, bold=True)
add_text(s, Inches(1.0), Inches(2.9), Inches(11), Inches(1.5),
         '① 我选的情境是 ____（等待/归途/邂逅）\n② 它要表达的情绪方向是 ____\n③ 我的依据（引用拉片观察：哪个动作/状态/关系让我这么判断）', size=16, color=GREY)
add_text(s, Inches(0.7), Inches(5.2), Inches(12), Inches(0.5),
         '同伴互评：1 次合作学习（引用 02_考题 3.1）', size=13, color=LIGHTGREY)
add_footer(s, 16)

# ============ P17 321出门条 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '321 出门条', '当场收 · 作为 W2 学情输入', CYAN)
add_card(s, Inches(0.7), Inches(1.7), Inches(5.8), Inches(1.8),
         '3 个收获', '今天最想记住的 3 点', BLUE, tsize=16, bsize=14)
add_card(s, Inches(6.9), Inches(1.7), Inches(5.8), Inches(1.8),
         '2 个疑问', '还没搞懂的 2 个问题', CYAN, tsize=16, bsize=14)
add_card(s, Inches(0.7), Inches(3.8), Inches(11.9), Inches(1.6),
         '1 个下次想学', '期待 W2 无对白创作的 1 个点', ORANGE, tsize=16, bsize=14)
add_text(s, Inches(0.7), Inches(5.6), Inches(12), Inches(0.5),
         '疑问条汇总进 W2 备课包（引用 02_考题 3.2）', size=13, color=LIGHTGREY)
add_footer(s, 17)

# ============ P18 总结金句 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s, RGBColor(0xE0, 0xF7, 0xFF), RGBColor(0xEF, 0xF6, 0xFF))
add_title_bar(s, '总结金句', '把今天的课带走', CORAL)
add_glow(s, Inches(9), Inches(2), Inches(4), Inches(4), CORAL, alpha=0.08)
add_text(s, Inches(0.7), Inches(2.3), Inches(11.5), Inches(0.8),
         '情境 = 种子', size=30, color=DARKBLUE, bold=True)
add_text(s, Inches(0.7), Inches(3.1), Inches(11.5), Inches(0.8),
         '画面 = 语言', size=30, color=BLUE, bold=True)
add_text(s, Inches(0.7), Inches(3.9), Inches(11.5), Inches(0.8),
         '人物动作 = 叙事', size=30, color=CYAN, bold=True)
add_text(s, Inches(0.7), Inches(5.3), Inches(11.5), Inches(0.6),
         '思政："情绪正，画面才拍得正"', size=16, color=ORANGE, bold=True)
add_footer(s, 18)

# ============ P19 下次预告 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '下次预告 · W2', '无对白创作（自创剧本+分镜）', BLUE)
add_card(s, Inches(0.7), Inches(1.8), Inches(5.8), Inches(2.6),
         'W2 你将做到', '自创无对白剧本（人物动作叙事）\n+ 分镜脚本（画面/景别/运镜/音效）\n+ AI 概念分镜初体验', CYAN, tsize=15, bsize=13)
add_card(s, Inches(6.9), Inches(1.8), Inches(5.8), Inches(2.6),
         '项目进度', 'P1 1/4 → 2/4\n画面点子 = 剧本骨架\n（教材①第三章核心）', ORANGE, tsize=15, bsize=13)
add_footer(s, 19)

# ============ P20 课后作业 ============
s = prs.slides.add_slide(blank)
add_gradient_bg(s)
add_title_bar(s, '课后作业', '为 W2 创作做准备', CYAN)
add_card(s, Inches(0.7), Inches(1.7), Inches(5.8), Inches(2.8),
         '阅读', '教材①第一/二章（通读）\n第三章（无对白练习，预习范例）', BLUE, tsize=15, bsize=13)
add_card(s, Inches(6.9), Inches(1.7), Inches(5.8), Inches(2.8),
         '完成', '情境理解笔记\n（选情境 + 情绪方向 + 画面点子雏形）', ORANGE, tsize=15, bsize=13)
add_text(s, Inches(0.7), Inches(4.9), Inches(12), Inches(0.6),
         '对齐授课计划第 1 行课外作业｜下一节交笔记，作为 W2 学情输入', size=13, color=LIGHTGREY)
add_footer(s, 20)

prs.save('05_课堂备课包/备课包_第1周_作品P1_情境认知与画面思维/06_PPT/PPT_W1_情境认知.pptx')
print('PPT_W1 已生成：20 页')
